"""Read Project Design Delivery Checklist workbooks.

The real sheet (template at `P:\\04 Voltas Templates and Schedules\\Templates`)
is one workbook per project: a header block (Project Number, Project Address,
Client Name / Number / Email) followed by a numbered table of delivery
stages — Intake through Sch CB — each with a Date and By column that gets
filled in as the project moves along.

Labels are located by text, not fixed cell addresses, so small layout edits
to the template don't break parsing.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

HEADER_LABELS = {
    "project number": "project_number",
    "project address": "address",
    "client name": "client_name",
    "client number": "client_number",
    "clients email": "client_email",
    "client email": "client_email",
}


@dataclass
class Stage:
    number: int
    activity: str
    done_on: date | None = None
    by: str = ""
    row: int = 0


@dataclass
class ProjectChecklist:
    source: str
    project_number: str = ""
    address: str = ""
    client_name: str = ""
    client_number: str = ""
    client_email: str = ""
    stages: list[Stage] = field(default_factory=list)

    @property
    def label(self) -> str:
        return self.project_number or self.address or Path(self.source).stem


def _clean(value: object) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%b %d, %Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def read_checklist(path: str | Path) -> ProjectChecklist:
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    ws = wb.active
    chk = ProjectChecklist(source=str(path))

    date_col = by_col = None
    in_stages = False
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [_clean(c) for c in (row or ())]
        if not any(cells):
            continue
        first = cells[0].lower().rstrip(":").strip()

        if not in_stages:
            if first in HEADER_LABELS:
                value = next((c for c in cells[1:] if c), "")
                setattr(chk, HEADER_LABELS[first], value)
            elif first.startswith("sl") and any("activity" in c.lower() for c in cells):
                lowered = [c.lower() for c in cells]
                date_col = next((j for j, c in enumerate(lowered) if c == "date"), 3)
                by_col = next((j for j, c in enumerate(lowered) if c == "by"), 4)
                in_stages = True
            continue

        # Stage rows: numeric Sl. No. plus an activity name. Anything else
        # (footer notes, the template-path line) ends the table.
        if not first.replace(".", "").isdigit():
            break
        activity = next((c for c in cells[1:] if c), "")
        if not activity:
            continue
        raw = list(row)
        chk.stages.append(
            Stage(
                number=int(float(first)),
                activity=activity,
                done_on=_as_date(raw[date_col]) if date_col < len(raw) else None,
                by=_clean(raw[by_col]) if by_col < len(raw) else "",
                row=i,
            )
        )
    return chk


def load_checklists(path: str | Path) -> list[ProjectChecklist]:
    """Load one checklist workbook, or every workbook in a folder."""
    path = Path(path)
    if path.is_dir():
        books = sorted(
            p
            for p in path.rglob("*.xls[xm]")
            if not p.name.startswith("~$") and not p.name.startswith(".")
        )
        return [read_checklist(p) for p in books]
    return [read_checklist(path)]
